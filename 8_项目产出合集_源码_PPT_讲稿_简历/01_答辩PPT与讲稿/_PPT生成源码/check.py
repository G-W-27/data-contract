# -*- coding: utf-8 -*-
"""QA：估算文本溢出、页面留白率与越界元素（无需 LibreOffice）。"""
import sys
from pptx import Presentation
from pptx.util import Emu

W_IN, H_IN = 13.3, 7.5
EMU = 914400.0


def units(s):
    cn = 0.0
    for ch in s:
        cn += 1.0 if ord(ch) > 0x2000 else 0.55
    return cn


def lines(text, w_in, pt):
    per = max(4, int(w_in / (pt / 72 * 1.06)))
    return max(1, -(-int(units(text) * 100) // (per * 100)))


def need_h(text, w_in, pt, line_factor=1.38):
    return lines(text, w_in, pt) * pt * line_factor / 72


def main(path):
    prs = Presentation(path)
    print("slides:", len(prs.slides))
    overflow, lowfill, outside = [], [], []
    for i, slide in enumerate(prs.slides, 1):
        area = 0.0
        for sh in slide.shapes:
            try:
                L, T = sh.left / EMU, sh.top / EMU
                w, h = sh.width / EMU, sh.height / EMU
            except Exception:
                continue
            if L < -0.02 or T < -0.02 or L + w > W_IN + 0.02 or T + h > H_IN + 0.02:
                outside.append((i, sh.shape_type, round(L, 2), round(T, 2), round(L + w, 2), round(T + h, 2)))
            area += w * h
            if not sh.has_text_frame:
                continue
            tf = sh.text_frame
            txt = tf.text
            if not txt.strip():
                continue
            # 取最大字号
            pt = 0.0
            for p in tf.paragraphs:
                for r in p.runs:
                    if r.font.size:
                        pt = max(pt, r.font.size.pt)
            if not pt:
                continue
            pad = 0.12
            nh = need_h(txt, max(w - 2 * pad, 0.3), pt)
            if nh > h + 0.06:
                overflow.append((i, round(pt, 1), round(nh, 2), round(h, 2), txt[:34].replace("\n", " ")))
        cover = area / (W_IN * H_IN)
        if cover < 0.42:
            lowfill.append((i, round(cover, 2)))
    print("\n== 可能溢出（需高 > 框高）==")
    for o in overflow[:40]:
        print(f"  p{o[0]:>3}  {o[1]}pt  need={o[2]}  box={o[3]}  {o[4]}")
    print(f"  共 {len(overflow)} 处")
    print("\n== 覆盖率偏低（<42%，可能留白明显）==")
    print("  " + ", ".join(f"p{a}:{b}" for a, b in lowfill))
    print("\n== 越界元素 ==")
    for o in outside[:20]:
        print("  ", o)
    print(f"  共 {len(outside)} 处")


if __name__ == "__main__":
    main(sys.argv[1])
